"""Export du rapport de production vers Excel, et reprise de l'historique de l'ancien classeur.

L'ancien classeur « Rapport de production » a une ligne par arrêt dans l'onglet « Datas » :
machine, jour (date et heure de saisie), équipe, type d'arrêt / perte, durée (min), commentaire.
"""
from dataclasses import dataclass, field
from datetime import datetime
from io import BytesIO

from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from inventaire.excel import normaliser
from produits.models import Machine

from . import calculs
from .models import Arret, CauseArret, Creneau, Equipe, Shift

TYPE_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
FEUILLE_HISTORIQUE = "Datas"
LIGNES_MAXI_AVANT_ENTETE = 30


class FichierIllisible(Exception):
    pass


# Export ------------------------------------------------------------------------------------------

def pourcentage(valeur):
    return None if valeur is None else round(float(valeur), 4)


def ecrire_feuille(feuille, titres, lignes, formats=None):
    feuille.append(titres)
    for position, titre in enumerate(titres, start=1):
        cellule = feuille.cell(row=1, column=position)
        cellule.font = Font(bold=True)
        cellule.fill = PatternFill("solid", start_color="DCE6F7")
        feuille.column_dimensions[get_column_letter(position)].width = max(10, len(titre) + 3)
    for ligne in lignes:
        feuille.append(ligne)
    for colonne, format_nombre in (formats or {}).items():
        for (cellule,) in feuille.iter_rows(min_row=2, min_col=colonne, max_col=colonne):
            cellule.number_format = format_nombre
    feuille.freeze_panes = "A2"


def exporter(shifts):
    """Deux feuilles : une ligne par shift avec ses totaux, puis une ligne par arrêt (comme l'onglet Datas)."""
    classeur = Workbook()
    feuille_shifts = classeur.active
    feuille_shifts.title = "Shifts"
    shifts = list(shifts)
    lignes = []
    for shift in shifts:
        bilan = shift.bilan()
        lignes.append([
            shift.jour, shift.creneau.nom, shift.equipe.code, shift.machine.code,
            ", ".join(element.troncon.produit.reference for element in bilan.troncons),
            bilan.theorique, bilan.arrets, bilan.marche, bilan.quantite,
            round(float(bilan.attendue)) if bilan.attendue else None, pourcentage(bilan.rendement),
            shift.matricule.numero if shift.matricule else "",
        ])
    ecrire_feuille(feuille_shifts, ["Date", "Créneau", "Équipe", "Machine", "Produits", "Temps théorique (min)",
                                    "Arrêts (min)", "Temps de marche (min)", "Quantité", "Quantité attendue",
                                    "Rendement", "Matricule"], lignes, {1: "DD/MM/YYYY", 11: "0%"})
    feuille_arrets = classeur.create_sheet("Arrêts")
    lignes = [
        [shift.machine.code, timezone.localtime(arret.saisi_le).replace(tzinfo=None), shift.equipe.code,
         shift.creneau.nom, arret.troncon.produit.reference if arret.troncon else "", arret.cause.libelle,
         arret.duree, arret.commentaire]
        for shift in shifts for arret in shift.arrets.all()
    ]
    ecrire_feuille(feuille_arrets, ["Machine", "Saisi le", "Équipe", "Créneau", "Produit", "Cause", "Durée (min)",
                                    "Commentaire"], lignes, {2: "DD/MM/YYYY HH:MM"})
    feuille_arrets.column_dimensions["H"].width = 50
    sortie = BytesIO()
    classeur.save(sortie)
    return sortie.getvalue()


# Reprise de l'historique ---------------------------------------------------------------------------

def colonne_de_l_entete(texte):
    """Clé de colonne d'un titre de l'onglet Datas, ou None. Les titres coupés (« Durée (r ») sont reconnus."""
    texte = normaliser(texte or "")
    if texte.startswith("machin"):
        return "machine"
    if texte.startswith("jour") or texte.startswith("date"):
        return "jour"
    if texte.startswith("equipe"):
        return "equipe"
    if texte.startswith("type") or texte.startswith("cause"):
        return "cause"
    if texte.startswith("duree"):
        return "duree"
    if texte.startswith("comment"):
        return "commentaire"
    return None


COLONNES_OBLIGATOIRES = ["machine", "jour", "equipe", "cause", "duree"]


def lire_date(valeur):
    if isinstance(valeur, datetime):
        return valeur
    texte = str(valeur or "").strip()
    for format_date in ("%d/%m/%y %H:%M", "%d/%m/%Y %H:%M", "%d/%m/%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(texte, format_date)
        except ValueError:
            continue
    return None


def lire_lignes(fichier, nom_feuille=FEUILLE_HISTORIQUE):
    """Lignes de l'onglet en dictionnaires, avec leur numéro de ligne Excel. Le titre peut précéder l'en-tête."""
    try:
        classeur = load_workbook(fichier, read_only=True, data_only=True)
    except Exception as erreur:
        raise FichierIllisible("Le fichier n'est pas un classeur Excel (.xlsx) lisible.") from erreur
    if nom_feuille not in classeur.sheetnames:
        raise FichierIllisible(f"Onglet « {nom_feuille} » absent (onglets : {', '.join(classeur.sheetnames)}).")
    lignes = classeur[nom_feuille].iter_rows(values_only=True)
    for numero, valeurs in enumerate(lignes, start=1):
        cles = [colonne_de_l_entete(str(valeur)) if valeur is not None else None for valeur in valeurs]
        if all(cle in cles for cle in COLONNES_OBLIGATOIRES):
            break
        if numero >= LIGNES_MAXI_AVANT_ENTETE:
            raise FichierIllisible("En-tête introuvable : il faut les colonnes Machine, Jour, Equipe, "
                                   "Type arrêt / perte et Durée.")
    else:
        raise FichierIllisible("L'onglet est vide.")
    resultat = []
    for numero_ligne, valeurs in enumerate(lignes, start=numero + 1):
        ligne = {cle: valeur for cle, valeur in zip(cles, valeurs) if cle}
        if any(valeur not in (None, "") for valeur in ligne.values()):
            resultat.append((numero_ligne, ligne))
    return resultat


@dataclass
class BilanImport:
    arrets_crees: int = 0
    doublons: int = 0
    shifts_crees: int = 0
    causes_creees: list = field(default_factory=list)
    equipes_creees: list = field(default_factory=list)
    ignorees: list = field(default_factory=list)  # (numéro de ligne, raison)


@transaction.atomic
def importer_historique(fichier, nom_feuille=FEUILLE_HISTORIQUE):
    """Reprend les arrêts de l'ancien classeur. Relancer l'import ne crée pas de doublons.

    Le shift d'un arrêt est retrouvé par l'heure de saisie : saisi à 12 h 19, l'arrêt va au créneau du matin ;
    à 2 h, à la nuit de la veille. Les machines inconnues (« Bande ») sont ignorées.
    """
    bilan = BilanImport()
    machines = {normaliser(machine.code): machine for machine in Machine.objects.all()}
    creneaux = list(Creneau.objects.all())
    causes = {normaliser(cause.libelle): cause for cause in CauseArret.objects.all()}
    equipes = {normaliser(equipe.code): equipe for equipe in Equipe.objects.all()}
    for numero, ligne in lire_lignes(fichier, nom_feuille):
        machine = machines.get(normaliser(str(ligne.get("machine") or "")))
        saisi_le = lire_date(ligne.get("jour"))
        texte_equipe = str(ligne.get("equipe") or "").strip()
        texte_cause = str(ligne.get("cause") or "").strip()
        try:
            duree = int(float(ligne.get("duree")))
        except (TypeError, ValueError):
            duree = 0
        if machine is None:
            bilan.ignorees.append((numero, f"machine inconnue « {ligne.get('machine') or ''} »"))
            continue
        if saisi_le is None:
            bilan.ignorees.append((numero, f"date illisible « {ligne.get('jour') or ''} »"))
            continue
        if not texte_equipe or not texte_cause or duree <= 0:
            bilan.ignorees.append((numero, "équipe, type d'arrêt ou durée manquant"))
            continue
        saisi_le = timezone.make_aware(saisi_le) if timezone.is_naive(saisi_le) else saisi_le
        jour, creneau = calculs.creneau_a_l_heure(saisi_le, creneaux)
        if creneau is None:
            bilan.ignorees.append((numero, f"saisi le {timezone.localtime(saisi_le):%d/%m/%Y à %H:%M}, hors de tout créneau"))
            continue
        equipe = equipes.get(normaliser(texte_equipe))
        if equipe is None:
            equipe = equipes[normaliser(texte_equipe)] = Equipe.objects.create(code=texte_equipe)
            bilan.equipes_creees.append(texte_equipe)
        cause = causes.get(normaliser(texte_cause))
        if cause is None:
            cause = causes[normaliser(texte_cause)] = CauseArret.objects.create(libelle=texte_cause)
            bilan.causes_creees.append(texte_cause)
        debut, fin = creneau.bornes(jour)
        shift, cree = Shift.objects.get_or_create(machine=machine, jour=jour, creneau=creneau,
                                                  defaults={"equipe": equipe, "debut": debut, "fin": fin,
                                                            "ouvert_le": saisi_le})
        bilan.shifts_crees += cree
        commentaire = str(ligne.get("commentaire") or "").strip()[:300]
        _, cree = Arret.objects.get_or_create(shift=shift, cause=cause, duree=duree, commentaire=commentaire,
                                              saisi_le=saisi_le)
        if cree:
            bilan.arrets_crees += 1
        else:
            bilan.doublons += 1
    return bilan

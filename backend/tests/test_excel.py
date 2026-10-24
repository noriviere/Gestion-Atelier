"""Import et export des articles au format Excel."""
from decimal import Decimal
from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import Workbook, load_workbook

from comptes import roles
from inventaire import excel
from inventaire.models import Article, Fournisseur, TypeInventaire

from .fabriques import creer_article, creer_type_inventaire, creer_utilisateur, lier_fournisseur

ENTETES = ["Référence", "Désignation", "Catégorie", "Origine", "N° de plan", "Type d'inventaire", "Emplacement", "Unité",
           "Stock", "Stock mini", "Mini de commande", "Fréquence d'inventaire (jours)", "Fournisseur préféré",
           "Réf. fournisseur", "Prix unitaire (€)", "Délai (jours)"]


def classeur(*lignes, entetes=ENTETES):
    """Fichier .xlsx en mémoire : une ligne d'en-têtes puis les lignes données."""
    document = Workbook()
    feuille = document.active
    feuille.append(entetes)
    for ligne in lignes:
        feuille.append(ligne)
    contenu = BytesIO()
    document.save(contenu)
    contenu.seek(0)
    return contenu


def ligne_patin(**remplacements):
    valeurs = dict(zip(ENTETES, ["PAT-X", "Patin taille X", "Pièce d'usure", "Design interne", "PL-1200 ind. A", "Magasin",
                                 "A1-09", "pièce", 12, 10, 20, 90, "Usinage Martin", "", 18.5, 10]))
    valeurs.update(remplacements)
    return [valeurs[entete] for entete in ENTETES]


class ExportTests(TestCase):
    def test_export_puis_relecture_identique(self):
        article = creer_article("BAN-40", stock=35, designation="Bande grain 40", categorie=Article.CONSOMMABLE)
        lien = lier_fournisseur(article, "Abrasifs", prefere=True, prix=Decimal("6.80"))
        lien.reference_fournisseur = "AI-B40"
        lien.save()
        contenu = excel.exporter_articles(Article.objects.all())
        feuille = load_workbook(BytesIO(contenu)).active
        self.assertEqual([cellule.value for cellule in feuille[1]], ENTETES)
        self.assertEqual([cellule.value for cellule in feuille[2]][:4], ["BAN-40", "Bande grain 40", "Consommable", "Article acheté"])
        analyse = excel.analyser_import(excel.lire_classeur(BytesIO(contenu)))
        self.assertEqual([ligne.action for ligne in analyse], ["inchange"])

    def test_modele_vide_n_a_que_les_entetes(self):
        feuille = load_workbook(BytesIO(excel.exporter_articles([]))).active
        self.assertEqual(feuille.max_row, 1)


class LectureTests(TestCase):
    def test_colonnes_dans_le_desordre_et_titres_approchants(self):
        lignes = excel.lire_classeur(classeur(["10", "PAT-X", "Patin", "usure", "acheté", "Magasin", "5"],
                                              entetes=["stock mini", "REF", "Designation", "categorie", "Origine",
                                                       "Type d'inventaire", "Quantité mini de commande"]))
        self.assertEqual(lignes[0]["valeurs"]["reference"], "PAT-X")
        self.assertEqual(lignes[0]["valeurs"]["stock_mini"], "10")
        self.assertEqual(lignes[0]["numero"], 2)

    def test_colonne_obligatoire_absente(self):
        with self.assertRaisesMessage(excel.FichierIllisible, "Stock mini"):
            excel.lire_classeur(classeur(entetes=[entete for entete in ENTETES if entete != "Stock mini"]))

    def test_fichier_qui_n_est_pas_un_classeur(self):
        with self.assertRaises(excel.FichierIllisible):
            excel.lire_classeur(BytesIO(b"pas un fichier excel"))

    def test_lignes_vides_ignorees(self):
        lignes = excel.lire_classeur(classeur(ligne_patin(), [None] * len(ENTETES), ligne_patin(**{"Référence": "PAT-Y"})))
        self.assertEqual([ligne["numero"] for ligne in lignes], [2, 4])


class AnalyseTests(TestCase):
    def analyser(self, *lignes):
        return excel.analyser_import(excel.lire_classeur(classeur(*lignes)))

    def test_nouvel_article(self):
        ligne, = self.analyser(ligne_patin())
        self.assertEqual(ligne.action, "creation")
        self.assertEqual(ligne.donnees["origine"], Article.DESIGN_INTERNE)
        self.assertEqual(ligne.donnees["stock"], Decimal("12"))
        self.assertIn("Type d'inventaire « Magasin » nouveau : il sera créé.", ligne.remarques)

    def test_erreurs_de_saisie(self):
        ligne, = self.analyser(ligne_patin(**{"Catégorie": "outil", "N° de plan": None, "Stock mini": "dix",
                                              "Fréquence d'inventaire (jours)": 7.5}))
        self.assertEqual(ligne.action, "erreur")
        self.assertEqual(len(ligne.erreurs), 4)

    def test_reference_en_double_dans_le_fichier(self):
        premiere, seconde = self.analyser(ligne_patin(), ligne_patin())
        self.assertEqual(premiere.action, "creation")
        self.assertEqual(seconde.erreurs, ["Référence déjà présente ligne 2."])

    def test_article_existant_modifie_sans_toucher_au_stock(self):
        creer_article("PAT-X", stock=3, stock_mini=10, mini_commande=20, designation="Patin taille X",
                      origine=Article.DESIGN_INTERNE, numero_plan="PL-1200 ind. A", emplacement="A1-09",
                      type_inventaire=creer_type_inventaire("Magasin"))
        ligne, = self.analyser(ligne_patin(**{"Stock mini": 8, "Fournisseur préféré": None, "Prix unitaire (€)": None,
                                              "Délai (jours)": None}))
        self.assertEqual(ligne.action, "modification")
        self.assertEqual(ligne.changements, ["Stock mini"])
        self.assertNotIn("stock", ligne.donnees)
        self.assertTrue(any("ignoré" in remarque for remarque in ligne.remarques))

    def test_article_achete_ignore_le_numero_de_plan(self):
        ligne, = self.analyser(ligne_patin(**{"Origine": "Article acheté", "Réf. fournisseur": "UM-PX"}))
        self.assertEqual(ligne.donnees["numero_plan"], "")
        self.assertEqual(ligne.fournisseur["reference_fournisseur"], "UM-PX")


class ApplicationTests(TestCase):
    def test_import_cree_article_stock_et_fournisseur_prefere(self):
        lier_fournisseur(creer_article("AUTRE"), "usinage martin")
        analyse = excel.analyser_import(excel.lire_classeur(classeur(ligne_patin())))
        self.assertEqual(excel.appliquer_import(analyse), (1, 0))
        article = Article.objects.get(reference="PAT-X")
        self.assertEqual(article.stock, Decimal("12"))
        self.assertEqual(article.mouvements.get().commentaire, "Stock initial (import Excel)")
        self.assertEqual(article.fournisseur_prefere.nom, "usinage martin")
        self.assertEqual(Fournisseur.objects.count(), 1)
        self.assertTrue(TypeInventaire.objects.filter(nom="Magasin").exists())

    def test_nouveau_fournisseur_prefere_remplace_l_ancien(self):
        article = creer_article("PAT-X", type_inventaire=creer_type_inventaire("Magasin"))
        lier_fournisseur(article, "Mécanique Sud", prefere=True)
        excel.appliquer_import(excel.analyser_import(excel.lire_classeur(classeur(ligne_patin()))))
        self.assertEqual(article.fournisseur_prefere.nom, "Usinage Martin")
        self.assertEqual(article.liens_fournisseurs.count(), 2)

    def test_lignes_en_erreur_non_importees(self):
        analyse = excel.analyser_import(excel.lire_classeur(classeur(ligne_patin(), ligne_patin(**{"Référence": "PAT-Y",
                                                                                                    "Désignation": None}))))
        self.assertEqual(excel.appliquer_import(analyse), (1, 0))
        self.assertFalse(Article.objects.filter(reference="PAT-Y").exists())


class PagesImportTests(TestCase):
    def test_apercu_puis_confirmation(self):
        self.client.force_login(creer_utilisateur("methode", roles.METHODE))
        fichier = SimpleUploadedFile("articles.xlsx", classeur(ligne_patin()).read())
        apercu = self.client.post("/articles/import/", {"fichier": fichier})
        self.assertContains(apercu, "Importer 1 article")
        self.assertFalse(Article.objects.exists())
        self.client.post("/articles/import/", {"etape": "confirmer"})
        self.assertTrue(Article.objects.filter(reference="PAT-X").exists())

    def test_confirmation_sans_apercu(self):
        self.client.force_login(creer_utilisateur("methode", roles.METHODE))
        self.assertRedirects(self.client.post("/articles/import/", {"etape": "confirmer"}), "/articles/import/")

    def test_operateur_exporte_mais_n_importe_pas(self):
        self.client.force_login(creer_utilisateur("operateur", roles.ATELIER))
        export = self.client.get("/articles/export.xlsx")
        self.assertEqual(export["Content-Type"], excel.TYPE_XLSX)
        self.assertEqual(self.client.get("/articles/import/").status_code, 403)

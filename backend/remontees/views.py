from io import BytesIO

from django.conf import settings
from django.contrib import messages
from django.db import transaction
from django.db.models import Count, Q
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.utils.text import slugify

from openpyxl import Workbook
from openpyxl.styles import Font

from comptes import roles
from inventaire.excel import TYPE_XLSX
from inventaire.models import Article, ArticleFournisseur, Mouvement
from produits.matrice import cle_naturelle
from produits.models import Machine

from . import regles, statistiques
from .forms import CorrectionForm, DemontageForm, MontageForm, ReglesForm, ThemesForm, valeurs_initiales_montage
from .models import FormatMatiere, HistoriqueRegle, Matricule, RegleMeule, SaisieMeule

NOMBRE_DERNIERES_SAISIES = 10
# Au-delà, la liste des saisies demande d'affiner les filtres (l'historique Excel compte des milliers de lignes).
NOMBRE_MAXIMUM_SAISIES = 300
NOMBRE_LIGNES_HISTORIQUE = 30


def machines_avec_meules_montees():
    """Machines dans l'ordre naturel (S2 avant S10), avec le nombre de meules montées et pas encore démontées."""
    machines = Machine.objects.annotate(meules_montees=Count(
        "saisies_meules", filter=Q(saisies_meules__demonte_le__isnull=True)))
    return sorted(machines, key=lambda machine: cle_naturelle(machine.code))


def remontees(request):
    """Écran commun des remontées : la machine d'abord, puis le sujet et ce qui attend sur cette machine."""
    machines = machines_avec_meules_montees()
    machine = next((machine for machine in machines if machine.code == request.GET.get("machine")), None)
    contexte = {"machines": machines, "machine_choisie": machine}
    if machine:
        saisies = SaisieMeule.objects.filter(machine=machine).select_related("meule", "format_matiere__produit")
        contexte["meules_montees"] = saisies.montees()
        contexte["dernieres_demontees"] = saisies.demontees().order_by("-demonte_le")[:NOMBRE_DERNIERES_SAISIES]
    return render(request, "remontees/remontees.html", contexte)


def contexte_formulaire(formulaire):
    """Ce que les gabarits de saisie ajoutent au formulaire : la liste des matricules."""
    return {"form": formulaire, "matricules": Matricule.objects.filter(actif=True)}


def table_formats():
    """Pour remontees.js : code lame, MA5, meules et machines prévues de chaque format et matière."""
    table = {}
    for ligne in FormatMatiere.objects.prefetch_related("meules", "machines"):
        table.setdefault(str(ligne.produit_id), {})[ligne.matiere] = {
            "code_lame": ligne.code_lame,
            "ma5": ligne.ma5,
            "meules": [str(meule.pk) for meule in ligne.meules.all()],
            "machines": [str(machine.pk) for machine in ligne.machines.all()],
        }
    return table


def matricule(numero):
    return Matricule.objects.get_or_create(numero=numero)[0]


def montage_meule(request):
    peut_sortir_du_stock = roles.a_le_droit(request.user, "stock_mouvements")
    machine = Machine.objects.filter(code=request.GET.get("machine")).first()
    formulaire = MontageForm(request.POST or None, initial=valeurs_initiales_montage(machine),
                             peut_sortir_du_stock=peut_sortir_du_stock)
    if request.method == "POST" and formulaire.is_valid():
        if formulaire.avertissements_a_confirmer:
            formulaire.marquer_avertissements_vus()
        else:
            saisie = enregistrer_montage(formulaire, request.user)
            message = f"Montage enregistré : meule {saisie.meule.reference} face {saisie.face} sur {saisie.machine.code}."
            if saisie.sortie_stock:
                message += f" Sortie de stock faite, nouveau stock {saisie.sortie_stock.stock_apres:g}."
            messages.success(request, message)
            return redirect(f"{reverse('remontees')}?machine={saisie.machine.code}")
    contexte = contexte_formulaire(formulaire)
    contexte.update({"machine": machine, "table_formats": table_formats(),
                     "controles": regles.controles_en_direct(regles.lire_reglages())})
    return render(request, "remontees/montage_meule.html", contexte)


@transaction.atomic
def enregistrer_montage(formulaire, utilisateur):
    donnees = formulaire.cleaned_data
    saisie = SaisieMeule.objects.create(
        machine=donnees["machine"], format_matiere=formulaire.format_matiere, face=donnees["face"],
        meule=donnees["meule"], lot_fournisseur=donnees["lot_fournisseur"], date_montage=donnees["date_montage"],
        ma5=donnees["ma5"] == FormatMatiere.MA5_OUI, pression_arrosage=donnees["pression_arrosage"],
        vitesse_rotation=donnees["vitesse_rotation"], vitesse_avance=donnees["vitesse_avance"],
        hauteur_initiale=donnees["hauteur_initiale"], compteur_debut=donnees["compteur_debut"],
        matricule_monteur=matricule(donnees["matricule_monteur"]),
        commentaire_montage=donnees["commentaire_montage"], monte_par=utilisateur,
        avertissements="\n".join(formulaire.avertissements),
    )
    if donnees.get("sortir_du_stock"):
        # Sortie ordinaire : les sorties de stock ne sont pas rattachées à une machine.
        saisie.sortie_stock = Mouvement.enregistrer(saisie.meule, Mouvement.SORTIE, 1, utilisateur=utilisateur,
                                                    commentaire=f"Montage de meule (remontée n° {saisie.pk})")
        saisie.save(update_fields=["sortie_stock"])
    return saisie


def demontage_meule(request, pk):
    saisie = get_object_or_404(SaisieMeule.objects.select_related("machine", "meule", "format_matiere__produit",
                                                                  "matricule_monteur"), pk=pk)
    if saisie.est_demontee:
        messages.info(request, "Cette meule est déjà démontée.")
        return redirect(saisie)
    formulaire = DemontageForm(request.POST or None, saisie=saisie)
    if request.method == "POST" and formulaire.is_valid():
        if formulaire.avertissements_a_confirmer:
            formulaire.marquer_avertissements_vus()
        else:
            enregistrer_demontage(saisie, formulaire, request.user)
            message = f"Démontage enregistré : {saisie.nombre_faces_brut} faces"
            if saisie.faces_par_meule_reference is not None:
                message += (f", soit {saisie.faces_par_meule_reference:.0f} faces pour une meule de "
                            f"{regles.HAUTEUR_REFERENCE_MM} mm")
            messages.success(request, message + ".")
            return redirect(saisie)
    contexte = contexte_formulaire(formulaire)
    contexte.update({"saisie": saisie, "aide_commentaire": regles.AIDE_COMMENTAIRE,
                     "seuil_non_finie": formulaire.reglages.seuil_non_finie,
                     "controles": regles.controles_en_direct(formulaire.reglages, saisie)})
    return render(request, "remontees/demontage_meule.html", contexte)


@transaction.atomic
def enregistrer_demontage(saisie, formulaire, utilisateur):
    donnees = formulaire.cleaned_data
    saisie.matricule_demonteur = matricule(donnees["matricule_demonteur"])
    saisie.compteur_fin = donnees["compteur_fin"]
    saisie.hauteur_finale = donnees["hauteur_finale"]
    saisie.probleme = donnees["probleme"]
    saisie.themes = donnees["themes"]
    saisie.commentaire = donnees["commentaire"]
    saisie.demonte_par = utilisateur
    saisie.demonte_le = timezone.now()
    saisie.marquer(formulaire.reglages)
    saisie.avertissements = "\n".join(filter(None, [saisie.avertissements, *formulaire.avertissements]))
    saisie.save()


def saisie_meule(request, pk):
    saisie = get_object_or_404(
        SaisieMeule.objects.select_related("machine", "meule", "format_matiere__produit", "matricule_monteur",
                                           "matricule_demonteur", "monte_par", "demonte_par", "corrige_par"), pk=pk)
    autres = list(SaisieMeule.objects.demontees().filter(meule=saisie.meule).exclude(pk=saisie.pk)
                  .select_related("format_matiere"))
    attendue = statistiques.mediane_attendue(autres, saisie.meule, saisie.machine, saisie.format_matiere.matiere,
                                             saisie.face)
    return render(request, "remontees/saisie_meule.html", {
        "saisie": saisie, "hauteur_reference": regles.HAUTEUR_REFERENCE_MM, "attendue": attendue,
        "minimum_fiable": statistiques.MINIMUM_FIABLE,
    })


# --- Consultation ----------------------------------------------------------------------------------
ETATS = [
    ("montee", "Montées, pas encore démontées"),
    ("demontee", "Démontées"),
    ("non_finie", "Non finies (X)"),
    ("defaut", "Défauts"),
    ("probleme", "Avec un problème signalé"),
    ("corrigee", "Corrigées"),
]


def etats_de_la_saisie(saisie):
    """Les états de la liste ETATS qui s'appliquent : sert au filtre de la démo statique (data-etat)."""
    etats = ["demontee" if saisie.est_demontee else "montee"]
    etats += [etat for etat, vrai in [("non_finie", saisie.non_finie), ("defaut", saisie.defaut),
                                      ("probleme", saisie.probleme), ("corrigee", saisie.corrige_le)] if vrai]
    return " ".join(etats)


def filtrer_saisies(saisies, filtres):
    if filtres["machine"]:
        saisies = saisies.filter(machine__code=filtres["machine"])
    if filtres["meule"].isdigit():
        saisies = saisies.filter(meule_id=filtres["meule"])
    if filtres["matiere"]:
        saisies = saisies.filter(format_matiere__matiere=filtres["matiere"])
    etat = filtres["etat"]
    if etat in ("montee", "demontee"):
        saisies = saisies.montees() if etat == "montee" else saisies.demontees()
    elif etat == "corrigee":
        saisies = saisies.filter(corrige_le__isnull=False)
    elif etat in ("non_finie", "defaut", "probleme"):
        saisies = saisies.filter(**{etat: True})
    if filtres["q"]:
        q = filtres["q"]
        saisies = saisies.filter(Q(lot_fournisseur__icontains=q) | Q(matricule_monteur__numero=q)
                                 | Q(matricule_demonteur__numero=q) | Q(commentaire__icontains=q)
                                 | Q(commentaire_montage__icontains=q) | Q(format_matiere__produit__reference__iexact=q))
    return saisies


def saisies_meules(request):
    """Toutes les saisies, filtrées : la page « Consultation » de l'ancienne application."""
    filtres = {cle: request.GET.get(cle, "").strip() for cle in ("q", "machine", "meule", "matiere", "etat")}
    saisies = SaisieMeule.objects.select_related("machine", "meule", "format_matiere__produit", "matricule_monteur",
                                                 "matricule_demonteur")
    if not settings.DEMO_STATIQUE:
        # Démo statique : toutes les lignes sont dans la page, demo.js masque celles hors filtre.
        saisies = filtrer_saisies(saisies, filtres)
    nombre = saisies.count()
    saisies = list(saisies[:NOMBRE_MAXIMUM_SAISIES] if not settings.DEMO_STATIQUE else saisies)
    for saisie in saisies:
        saisie.etats = etats_de_la_saisie(saisie)
    return render(request, "remontees/saisies_meules.html", {
        "saisies": saisies, "nombre": nombre, "tronquee": nombre > len(saisies), "filtres": filtres,
        "machines": sorted(Machine.objects.filter(saisies_meules__isnull=False).distinct(),
                           key=lambda machine: cle_naturelle(machine.code)),
        "meules": Article.objects.filter(saisies_meules__isnull=False).distinct().order_by("reference"),
        "matieres": FormatMatiere.MATIERES, "etats": ETATS,
    })


def saisies_demontees():
    """Les saisies démontées avec tout ce que les statistiques lisent, en une requête (plus les prévisions)."""
    return list(SaisieMeule.objects.demontees().select_related("machine", "meule", "format_matiere__produit")
                .prefetch_related("format_matiere__machines", "format_matiere__meules"))


def fournisseurs_preferes(articles):
    """{article_id: nom du fournisseur préféré} pour les meules des saisies."""
    return dict(ArticleFournisseur.objects.filter(article__in=articles, prefere=True)
                .values_list("article_id", "fournisseur__nom"))


def statistiques_meules(request):
    saisies = saisies_demontees()
    reglages = regles.lire_reglages()
    prix = reglages.valeur("prix_meule")
    annee = timezone.localdate().year
    fournisseurs = fournisseurs_preferes({saisie.meule_id for saisie in saisies})
    return render(request, "remontees/statistiques_meules.html", {
        "lignes": statistiques.lignes_pour_les_statistiques(saisies, fournisseurs),
        "pertes": [statistiques.perte(saisies, annee, prix), statistiques.perte(saisies, annee - 1, prix)],
        "coefficient_eloigne": float(reglages.valeur("duree_eloignee", "coefficient")),
        "hauteur_reference": regles.HAUTEUR_REFERENCE_MM,
    })


def controle_meules(request):
    anomalies = statistiques.anomalies(saisies_demontees(), regles.lire_reglages())
    gravite = request.GET.get("gravite", "")
    nombres = {gravite_possible: sum(1 for anomalie in anomalies if anomalie["gravite"] == gravite_possible)
               for gravite_possible in statistiques.GRAVITES}
    recherche = request.GET.get("q", "").strip()
    if not settings.DEMO_STATIQUE:
        # Démo statique : toutes les lignes sont dans la page, demo.js masque celles hors filtre.
        if gravite:
            anomalies = [anomalie for anomalie in anomalies if slugify(anomalie["gravite"]) == gravite]
        if recherche:
            anomalies = [anomalie for anomalie in anomalies if recherche.lower() in
                         f"{anomalie['champ']} {anomalie['motif']} {anomalie['valeur']}".lower()]
    return render(request, "remontees/controle_meules.html", {
        "anomalies": anomalies, "gravites": [(gravite_possible, nombres[gravite_possible])
                                             for gravite_possible in statistiques.GRAVITES],
        "gravite_choisie": gravite, "recherche": recherche, "grave": statistiques.GRAVE, "verif": statistiques.VERIF,
    })


COLONNES_CONTROLE = ["Gravité", "Saisie n°", "Date de montage", "Machine", "Meule", "Format", "Face", "Champ", "Valeur",
                     "Motif"]


def export_controle_meules(request):
    """Le contrôle qualité en Excel : une ligne par anomalie, pour corriger ou expliquer hors de l'application."""
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Contrôle qualité"
    feuille.append(COLONNES_CONTROLE)
    for cellule in feuille[1]:
        cellule.font = Font(bold=True)
    for anomalie in statistiques.anomalies(saisies_demontees(), regles.lire_reglages()):
        saisie = anomalie["saisie"]
        feuille.append([anomalie["gravite"], saisie.pk, saisie.date_montage, saisie.machine.code,
                        saisie.meule.reference, str(saisie.format_matiere), saisie.face, anomalie["champ"],
                        anomalie["valeur"], anomalie["motif"]])
        feuille.cell(feuille.max_row, 3).number_format = "DD/MM/YYYY"
    for colonne, largeur in zip("ABCDEFGHIJ", [16, 10, 14, 10, 14, 14, 6, 28, 14, 60]):
        feuille.column_dimensions[colonne].width = largeur
    feuille.freeze_panes = "A2"
    contenu = BytesIO()
    classeur.save(contenu)
    reponse = HttpResponse(contenu.getvalue(), content_type=TYPE_XLSX)
    reponse["Content-Disposition"] = f'attachment; filename="controle-meules-{timezone.localdate():%Y-%m-%d}.xlsx"'
    return reponse


# --- Correction (Méthode et management) ------------------------------------------------------------
def modifier_saisie_meule(request, pk):
    saisie = get_object_or_404(SaisieMeule.objects.select_related("machine", "meule", "format_matiere__produit",
                                                                  "matricule_monteur", "matricule_demonteur"), pk=pk)
    formulaire = CorrectionForm(request.POST or None, instance=saisie)
    if request.method == "POST" and formulaire.is_valid():
        enregistrer_correction(formulaire, request.user)
        messages.success(request, "Saisie corrigée. Les valeurs calculées et les marques X et Défaut sont recalculées.")
        return redirect(saisie)
    return render(request, "remontees/modifier_saisie_meule.html", {"form": formulaire, "saisie": saisie,
                                                                     "matricules": Matricule.objects.filter(actif=True)})


@transaction.atomic
def enregistrer_correction(formulaire, utilisateur):
    saisie = formulaire.save(commit=False)
    donnees = formulaire.cleaned_data
    saisie.matricule_monteur = matricule(donnees["matricule_monteur"])
    if saisie.est_demontee:
        saisie.matricule_demonteur = matricule(donnees["matricule_demonteur"])
        saisie.marquer(regles.lire_reglages())
    saisie.corrige_par = utilisateur
    saisie.corrige_le = timezone.now()
    saisie.save()


def supprimer_saisie_meule(request, pk):
    saisie = get_object_or_404(SaisieMeule.objects.select_related("meule", "machine"), pk=pk)
    if request.method != "POST":
        return redirect(saisie)
    texte = f"Saisie n° {saisie.pk} supprimée (meule {saisie.meule.reference} sur {saisie.machine.code})."
    if saisie.sortie_stock_id:
        texte += " La sortie de stock du montage reste : si la meule n'a pas servi, faites une entrée de stock."
    saisie.delete()
    messages.success(request, texte)
    return redirect("saisies_meules")


# --- Back-office des règles ------------------------------------------------------------------------
def historiser(phrases, utilisateur):
    HistoriqueRegle.objects.bulk_create([HistoriqueRegle(utilisateur=utilisateur, texte=phrase[:300])
                                         for phrase in phrases])


def regles_meules(request):
    reglages = regles.lire_reglages()
    formulaire_envoye = request.POST.get("formulaire") if request.method == "POST" else None
    form_regles = ReglesForm(request.POST if formulaire_envoye == "regles" else None, reglages=reglages)
    form_themes = ThemesForm(request.POST if formulaire_envoye == "themes" else None)
    if formulaire_envoye == "origine":
        with transaction.atomic():
            RegleMeule.objects.all().delete()
            historiser(["Toutes les règles remises aux valeurs d'origine"], request.user)
        messages.success(request, "Règles remises aux valeurs d'origine. Les thèmes ne changent pas.")
        return redirect("regles_meules")
    if formulaire_envoye == "regles" and form_regles.is_valid():
        nouveaux, phrases = form_regles.changements()
        with transaction.atomic():
            for code, (actif, valeurs) in nouveaux.items():
                RegleMeule.objects.update_or_create(code=code, defaults={"actif": actif, "valeurs": valeurs})
            historiser(phrases, request.user)
        messages.success(request, f"Règles enregistrées ({len(phrases)} changement{'s' if len(phrases) > 1 else ''}).")
        return redirect("regles_meules")
    if formulaire_envoye == "themes" and form_themes.is_valid():
        with transaction.atomic():
            phrases = form_themes.enregistrer()
            historiser(phrases, request.user)
        messages.success(request, f"Thèmes enregistrés ({len(phrases)} changement{'s' if len(phrases) > 1 else ''}).")
        return redirect("regles_meules")
    return render(request, "remontees/regles_meules.html", {
        "form_regles": form_regles, "form_themes": form_themes, "theme_defaut": regles.THEME_DEFAUT,
        "historique": HistoriqueRegle.objects.select_related("utilisateur")[:NOMBRE_LIGNES_HISTORIQUE],
    })

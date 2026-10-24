"""Import et export des articles au format Excel (.xlsx).

Le même tableau sert aux deux sens : un export peut être corrigé puis réimporté.
Une ligne est rattachée à un article existant par sa référence ; sinon l'article est créé.
"""
import unicodedata
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .models import Article, ArticleFournisseur, Fournisseur, Mouvement, TypeInventaire

COLONNES = [
    # clé, titre de colonne, obligatoire, explication
    ("reference", "Référence", True, "Référence interne de l'article, unique. Sert à retrouver l'article à la réimportation."),
    ("designation", "Désignation", True, ""),
    ("categorie", "Catégorie", True, "Pièce d'usure ou Consommable."),
    ("origine", "Origine", True, "Article acheté ou Design interne."),
    ("numero_plan", "N° de plan", False, "Obligatoire pour un design interne, ignoré pour un article acheté."),
    ("type_inventaire", "Type d'inventaire", True, "Créé s'il n'existe pas encore."),
    ("emplacement", "Emplacement", False, ""),
    ("unite", "Unité", False, "pièce si vide."),
    ("stock", "Stock", False, "Pris en compte à la création seulement : ensuite, le stock change par les mouvements et les comptages."),
    ("stock_mini", "Stock mini", True, ""),
    ("mini_commande", "Mini de commande", True, "Quantité minimale par commande."),
    ("frequence_inventaire", "Fréquence d'inventaire (jours)", False, "90 si vide."),
    ("fournisseur", "Fournisseur préféré", False, "Créé s'il n'existe pas encore. Les autres fournisseurs se gèrent dans la fiche article."),
    ("reference_fournisseur", "Réf. fournisseur", False, "Pour un article acheté."),
    ("prix", "Prix unitaire (€)", False, ""),
    ("delai", "Délai (jours)", False, ""),
]
TITRES = {cle: titre for cle, titre, _, _ in COLONNES}
LIBELLES_CATEGORIES = dict(Article.CATEGORIES)
LIBELLES_ORIGINES = dict(Article.ORIGINES)
TYPE_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class FichierIllisible(Exception):
    pass


def normaliser(texte):
    """« Pièce d'usure » -> « piecedusure » : compare des saisies sans tenir compte des accents, espaces et majuscules."""
    sans_accents = unicodedata.normalize("NFKD", str(texte)).encode("ascii", "ignore").decode()
    return "".join(caractere for caractere in sans_accents.lower() if caractere.isalnum())


COLONNE_PAR_TITRE = {normaliser(titre): cle for cle, titre, _, _ in COLONNES}
COLONNE_PAR_TITRE.update({normaliser(alias): cle for alias, cle in [
    ("Réf", "reference"), ("Référence article", "reference"), ("Numéro de plan", "numero_plan"), ("Plan", "numero_plan"),
    ("Fournisseur", "fournisseur"), ("Prix", "prix"), ("Délai", "delai"), ("Fréquence d'inventaire", "frequence_inventaire"),
    ("Quantité mini de commande", "mini_commande"),
]})
CATEGORIE_PAR_SAISIE = {normaliser(saisie): code for code, libelle in Article.CATEGORIES for saisie in (code, libelle)}
CATEGORIE_PAR_SAISIE.update({normaliser("usure"): Article.PIECE_USURE, normaliser("consommable"): Article.CONSOMMABLE})
ORIGINE_PAR_SAISIE = {normaliser(saisie): code for code, libelle in Article.ORIGINES for saisie in (code, libelle)}
ORIGINE_PAR_SAISIE.update({normaliser("acheté"): Article.ACHAT, normaliser("achat"): Article.ACHAT,
                           normaliser("interne"): Article.DESIGN_INTERNE})


# Export ------------------------------------------------------------------------------------------

def valeurs_export(article):
    lien = next((lien for lien in article.liens_fournisseurs.all() if lien.prefere), None)
    return {
        "reference": article.reference,
        "designation": article.designation,
        "categorie": LIBELLES_CATEGORIES[article.categorie],
        "origine": LIBELLES_ORIGINES[article.origine],
        "numero_plan": article.numero_plan,
        "type_inventaire": article.type_inventaire.nom,
        "emplacement": article.emplacement,
        "unite": article.unite,
        "stock": article.stock,
        "stock_mini": article.stock_mini,
        "mini_commande": article.mini_commande,
        "frequence_inventaire": article.frequence_inventaire,
        "fournisseur": lien.fournisseur.nom if lien else "",
        "reference_fournisseur": lien.reference_fournisseur if lien else "",
        "prix": lien.prix if lien else None,
        "delai": lien.delai if lien else None,
    }


def exporter_articles(articles):
    """Classeur .xlsx (en octets) : une feuille Articles, une feuille Mode d'emploi."""
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Articles"
    feuille.append([titre for _, titre, _, _ in COLONNES])
    for cellule, (_, _, obligatoire, _) in zip(feuille[1], COLONNES):
        cellule.font = Font(bold=True)
        if obligatoire:
            cellule.fill = PatternFill("solid", start_color="DCE6F7")
    for article in articles:
        valeurs = valeurs_export(article)
        feuille.append([float(valeur) if isinstance(valeur, Decimal) else valeur for valeur in
                        (valeurs[cle] for cle, _, _, _ in COLONNES)])
    for position, (_, titre, _, _) in enumerate(COLONNES, start=1):
        feuille.column_dimensions[get_column_letter(position)].width = max(12, len(titre) + 4)
    feuille.column_dimensions["B"].width = 36
    feuille.freeze_panes = "B2"
    for cle, choix in (("categorie", LIBELLES_CATEGORIES.values()), ("origine", LIBELLES_ORIGINES.values())):
        lettre = get_column_letter([cle_colonne for cle_colonne, _, _, _ in COLONNES].index(cle) + 1)
        liste = DataValidation(type="list", formula1='"' + ",".join(choix) + '"', allow_blank=False)
        liste.add(f"{lettre}2:{lettre}2000")
        feuille.add_data_validation(liste)

    mode_emploi = classeur.create_sheet("Mode d'emploi")
    mode_emploi.append(["Colonne", "Obligatoire", "Explication"])
    for cellule in mode_emploi[1]:
        cellule.font = Font(bold=True)
    for _, titre, obligatoire, explication in COLONNES:
        mode_emploi.append([titre, "oui" if obligatoire else "", explication])
    mode_emploi.append([])
    mode_emploi.append(["Une ligne dont la référence existe déjà modifie l'article ; sinon l'article est créé."])
    mode_emploi.append(["Les colonnes en bleu sont obligatoires. L'ordre des colonnes n'a pas d'importance."])
    mode_emploi.column_dimensions["A"].width = 30
    mode_emploi.column_dimensions["B"].width = 12
    mode_emploi.column_dimensions["C"].width = 100

    sortie = BytesIO()
    classeur.save(sortie)
    return sortie.getvalue()


# Lecture -----------------------------------------------------------------------------------------

def texte_de_cellule(valeur):
    """Valeur de cellule en texte : 14.0 -> « 14 », vide -> « »."""
    if valeur is None:
        return ""
    if isinstance(valeur, float) and valeur.is_integer():
        return str(int(valeur))
    return str(valeur).strip()


def lire_classeur(fichier):
    """Lignes de la première feuille, en dictionnaires {clé de colonne: texte}, avec leur numéro de ligne Excel."""
    try:
        classeur = load_workbook(fichier, read_only=True, data_only=True)
    except Exception as erreur:
        raise FichierIllisible("Le fichier n'est pas un classeur Excel (.xlsx) lisible.") from erreur
    feuille = classeur.worksheets[0]
    lignes = feuille.iter_rows(values_only=True)
    entetes = next(lignes, None)
    if not entetes:
        raise FichierIllisible("La première feuille est vide.")
    cles = [COLONNE_PAR_TITRE.get(normaliser(entete)) if entete is not None else None for entete in entetes]
    manquantes = [TITRES[cle] for cle, _, obligatoire, _ in COLONNES if obligatoire and cle not in cles]
    if manquantes:
        raise FichierIllisible(f"Colonnes obligatoires absentes : {', '.join(manquantes)}.")
    resultat = []
    for numero, valeurs in enumerate(lignes, start=2):
        ligne = {cle: texte_de_cellule(valeur) for cle, valeur in zip(cles, valeurs) if cle}
        if any(ligne.values()):
            resultat.append({"numero": numero, "valeurs": ligne})
    return resultat


# Analyse et import -------------------------------------------------------------------------------

@dataclass
class LigneImport:
    numero: int
    reference: str
    article: Article = None
    donnees: dict = field(default_factory=dict)
    fournisseur: dict = field(default_factory=dict)
    erreurs: list = field(default_factory=list)
    remarques: list = field(default_factory=list)
    changements: list = field(default_factory=list)

    @property
    def action(self):
        if self.erreurs:
            return "erreur"
        if self.article is None:
            return "creation"
        return "modification" if self.changements else "inchange"


# Les champs numériques de l'article ont au plus 10 chiffres dont 2 après la virgule.
NOMBRE_MAXI = Decimal("99999999")


def lire_nombre(texte, titre, erreurs, entier=False, obligatoire=False):
    if not texte:
        if obligatoire:
            erreurs.append(f"{titre} manquant.")
        return None
    try:
        nombre = Decimal(texte.replace(",", ".").replace(" ", "").replace(" ", ""))
    except InvalidOperation:
        erreurs.append(f"{titre} illisible : « {texte} ».")
        return None
    if not nombre.is_finite() or nombre > NOMBRE_MAXI:
        erreurs.append(f"{titre} illisible : « {texte} ».")
        return None
    if nombre < 0:
        erreurs.append(f"{titre} négatif : « {texte} ».")
        return None
    if entier:
        if nombre != nombre.to_integral_value():
            erreurs.append(f"{titre} doit être un nombre entier : « {texte} ».")
            return None
        return int(nombre)
    return nombre


def trouver_ou_creer(modele, nom):
    """Fournisseur ou type d'inventaire de ce nom, sans tenir compte des majuscules ; créé s'il n'existe pas."""
    return modele.objects.filter(nom__iexact=nom).first() or modele.objects.create(nom=nom)


def verifier_longueurs(ligne):
    """Un texte plus long que sa colonne en base ferait échouer tout l'import à l'enregistrement."""
    # (colonne du fichier, modèle, champ en base, texte)
    textes = [("reference", Article, "reference", ligne.reference),
              ("type_inventaire", TypeInventaire, "nom", ligne.donnees["type_inventaire"])]
    textes += [(cle, Article, cle, ligne.donnees[cle]) for cle in ("designation", "numero_plan", "emplacement", "unite")]
    if ligne.fournisseur:
        textes += [("fournisseur", Fournisseur, "nom", ligne.fournisseur["nom"]),
                   ("reference_fournisseur", ArticleFournisseur, "reference_fournisseur",
                    ligne.fournisseur["reference_fournisseur"])]
    for cle, modele, champ, texte in textes:
        longueur_maxi = modele._meta.get_field(champ).max_length
        if len(texte) > longueur_maxi:
            ligne.erreurs.append(f"{TITRES.get(cle, cle)} trop long : {len(texte)} caractères pour {longueur_maxi} au plus.")


def analyser_ligne(numero, valeurs, articles_existants):
    reference = valeurs.get("reference", "")
    ligne = LigneImport(numero=numero, reference=reference, article=articles_existants.get(reference))
    erreurs = ligne.erreurs
    if not reference:
        erreurs.append("Référence manquante.")
    designation = valeurs.get("designation", "")
    if not designation:
        erreurs.append("Désignation manquante.")
    categorie = CATEGORIE_PAR_SAISIE.get(normaliser(valeurs.get("categorie", "")))
    if categorie is None:
        erreurs.append(f"Catégorie inconnue : « {valeurs.get('categorie', '')} » (Pièce d'usure ou Consommable).")
    origine = ORIGINE_PAR_SAISIE.get(normaliser(valeurs.get("origine", "")))
    if origine is None:
        erreurs.append(f"Origine inconnue : « {valeurs.get('origine', '')} » (Article acheté ou Design interne).")
    numero_plan = valeurs.get("numero_plan", "") if origine == Article.DESIGN_INTERNE else ""
    if origine == Article.DESIGN_INTERNE and not numero_plan:
        erreurs.append("Design interne sans numéro de plan.")
    type_inventaire = valeurs.get("type_inventaire", "")
    if not type_inventaire:
        erreurs.append("Type d'inventaire manquant.")
    stock = lire_nombre(valeurs.get("stock", ""), "Stock", erreurs)
    frequence = lire_nombre(valeurs.get("frequence_inventaire", ""), "Fréquence d'inventaire", erreurs, entier=True)
    ligne.donnees = {
        "designation": designation,
        "categorie": categorie,
        "origine": origine,
        "numero_plan": numero_plan,
        "type_inventaire": type_inventaire,
        "emplacement": valeurs.get("emplacement", ""),
        "unite": valeurs.get("unite", "") or "pièce",
        "stock_mini": lire_nombre(valeurs.get("stock_mini", ""), "Stock mini", erreurs, obligatoire=True),
        "mini_commande": lire_nombre(valeurs.get("mini_commande", ""), "Mini de commande", erreurs, obligatoire=True),
        "frequence_inventaire": frequence if frequence is not None else 90,
    }
    if ligne.article is None:
        ligne.donnees["stock"] = stock or Decimal(0)
    elif stock is not None and stock != ligne.article.stock:
        ligne.remarques.append(f"Stock du fichier ({stock.normalize():f}) ignoré : le stock d'un article existant "
                               "se corrige par un comptage.")
    if valeurs.get("fournisseur"):
        ligne.fournisseur = {
            "nom": valeurs["fournisseur"],
            "reference_fournisseur": valeurs.get("reference_fournisseur", "") if origine != Article.DESIGN_INTERNE else "",
            "prix": lire_nombre(valeurs.get("prix", ""), "Prix", erreurs),
            "delai": lire_nombre(valeurs.get("delai", ""), "Délai", erreurs, entier=True),
        }
    elif any(valeurs.get(cle) for cle in ("reference_fournisseur", "prix", "delai")):
        erreurs.append("Réf. fournisseur, prix ou délai sans fournisseur préféré.")
    verifier_longueurs(ligne)
    if ligne.article is not None and not erreurs:
        ligne.changements = changements(ligne.article, ligne.donnees, ligne.fournisseur)
    return ligne


def changements(article, donnees, fournisseur):
    """Titres des colonnes dont la valeur du fichier diffère de l'article en base."""
    actuelles = valeurs_export(article)
    nouvelles = {
        **donnees,
        "categorie": LIBELLES_CATEGORIES[donnees["categorie"]],
        "origine": LIBELLES_ORIGINES[donnees["origine"]],
    }
    if fournisseur:
        nouvelles.update({"fournisseur": fournisseur["nom"], "reference_fournisseur": fournisseur["reference_fournisseur"]})
        nouvelles.update({cle: fournisseur[cle] for cle in ("prix", "delai") if fournisseur[cle] is not None})
    return [TITRES[cle] for cle, valeur in nouvelles.items() if cle in actuelles and cle != "stock" and valeur != actuelles[cle]]


def analyser_import(lignes_lues):
    """Contrôle chaque ligne sans rien enregistrer ; signale aussi les références en double dans le fichier."""
    references = [ligne["valeurs"].get("reference", "") for ligne in lignes_lues]
    noms_existants = {
        "type_inventaire": {nom.lower() for nom in TypeInventaire.objects.values_list("nom", flat=True)},
        "fournisseur": {nom.lower() for nom in Fournisseur.objects.values_list("nom", flat=True)},
    }
    articles_existants = {
        article.reference: article
        for article in Article.objects.filter(reference__in=references)
        .select_related("type_inventaire").prefetch_related("liens_fournisseurs__fournisseur")
    }
    analyse = []
    deja_vues = {}
    for ligne_lue in lignes_lues:
        ligne = analyser_ligne(ligne_lue["numero"], ligne_lue["valeurs"], articles_existants)
        if ligne.reference and ligne.reference in deja_vues:
            ligne.erreurs.append(f"Référence déjà présente ligne {deja_vues[ligne.reference]}.")
        deja_vues.setdefault(ligne.reference, ligne.numero)
        for cle, libelle in (("type_inventaire", "Type d'inventaire"), ("fournisseur", "Fournisseur")):
            nom = ligne_lue["valeurs"].get(cle, "")
            if nom and nom.lower() not in noms_existants[cle]:
                ligne.remarques.append(f"{libelle} « {nom} » nouveau : il sera créé.")
        analyse.append(ligne)
    return analyse


@transaction.atomic
def appliquer_import(analyse, utilisateur=None):
    """Enregistre les lignes sans erreur ; renvoie le nombre d'articles créés et modifiés."""
    nombre_crees = nombre_modifies = 0
    for ligne in analyse:
        if ligne.action in ("erreur", "inchange"):
            continue
        donnees = dict(ligne.donnees)
        donnees["type_inventaire"] = trouver_ou_creer(TypeInventaire, donnees["type_inventaire"])
        stock_initial = donnees.pop("stock", None)
        if ligne.article is None:
            article = Article.objects.create(reference=ligne.reference, **donnees)
            if stock_initial:
                Mouvement.enregistrer(article, Mouvement.ENTREE, stock_initial, utilisateur=utilisateur,
                                      commentaire="Stock initial (import Excel)")
            nombre_crees += 1
        else:
            article = ligne.article
            for champ, valeur in donnees.items():
                setattr(article, champ, valeur)
            # Sans toucher au stock : un mouvement saisi pendant l'import ne doit pas être écrasé.
            article.save(update_fields=list(donnees))
            nombre_modifies += 1
        if ligne.fournisseur:
            definir_fournisseur_prefere(article, ligne.fournisseur)
    return nombre_crees, nombre_modifies


def definir_fournisseur_prefere(article, fournisseur):
    nouveau_prefere = trouver_ou_creer(Fournisseur, fournisseur["nom"])
    article.liens_fournisseurs.exclude(fournisseur=nouveau_prefere).update(prefere=False)
    lien, _ = ArticleFournisseur.objects.get_or_create(article=article, fournisseur=nouveau_prefere)
    lien.prefere = True
    lien.reference_fournisseur = fournisseur["reference_fournisseur"]
    if fournisseur["prix"] is not None:
        lien.prix = fournisseur["prix"]
    if fournisseur["delai"] is not None:
        lien.delai = fournisseur["delai"]
    lien.save()

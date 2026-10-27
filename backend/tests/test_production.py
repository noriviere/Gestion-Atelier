"""Rapport de production : créneaux, calculs, saisie du shift, droits, rapport filtrable et reprise de l'Excel."""
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from io import BytesIO

from django.test import SimpleTestCase, TestCase
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from comptes import roles
from production import calculs
from production.excel import importer_historique
from production.models import Arret, CauseArret, Correction, Creneau, Equipe, Shift, Troncon
from produits.models import Machine, Produit, ProduitMachine
from remontees.models import Matricule

from .fabriques import creer_utilisateur

LUNDI = date(2026, 10, 19)


def a(jour, heure, minute=0):
    return timezone.make_aware(datetime.combine(jour, time(heure, minute)))


def creneau(nom):
    return Creneau.objects.get(nom=nom)


class CreneauxTests(TestCase):
    """Les créneaux de départ : matin 5-13 h, après-midi 13-21 h, nuit 21-5 h, weekend 6-18 h."""

    def test_listes_de_depart(self):
        self.assertEqual(list(Equipe.objects.values_list("code", flat=True)), ["E1", "E2", "E3", "E4"])
        self.assertEqual([c.duree for c in Creneau.objects.all()], [480, 480, 480, 720])
        self.assertTrue(CauseArret.objects.get(libelle="Autre").commentaire_obligatoire)
        self.assertTrue(CauseArret.objects.filter(libelle="Outillages").exists())

    def test_la_nuit_finit_le_lendemain(self):
        debut, fin = creneau("Nuit").bornes(LUNDI)
        self.assertEqual((debut, fin), (a(LUNDI, 21), a(LUNDI + timedelta(days=1), 5)))

    def test_creneau_selon_l_heure(self):
        creneaux = list(Creneau.objects.all())
        mardi, samedi = LUNDI + timedelta(days=1), LUNDI + timedelta(days=5)
        self.assertEqual(calculs.creneau_a_l_heure(a(LUNDI, 12, 19), creneaux), (LUNDI, creneau("Matin")))
        self.assertEqual(calculs.creneau_a_l_heure(a(LUNDI, 13), creneaux), (LUNDI, creneau("Après-midi")))
        # À 2 h du matin, c'est la nuit qui a commencé la veille.
        self.assertEqual(calculs.creneau_a_l_heure(a(mardi, 2), creneaux), (LUNDI, creneau("Nuit")))
        self.assertEqual(calculs.creneau_a_l_heure(a(samedi, 10), creneaux), (samedi, creneau("Weekend")))
        # Samedi 20 h : entre deux créneaux.
        self.assertEqual(calculs.creneau_a_l_heure(a(samedi, 20), creneaux), (None, None))


class CalculsTests(SimpleTestCase):
    def test_quantite_avec_ou_sans_remise_a_zero(self):
        self.assertEqual(calculs.quantite(0, 1450), 1450)
        self.assertEqual(calculs.quantite(None, 1450), 1450)
        self.assertEqual(calculs.quantite(12000, 13450), 1450)
        self.assertIsNone(calculs.quantite(0, None))

    def test_rendement(self):
        self.assertEqual(calculs.quantite_attendue(Decimal("200"), 240), Decimal(800))
        self.assertEqual(calculs.rendement(720, Decimal(800)), Decimal("0.9"))
        self.assertIsNone(calculs.quantite_attendue(None, 240))
        self.assertIsNone(calculs.rendement(720, None))
        self.assertIsNone(calculs.rendement(None, Decimal(800)))


class FabriqueShift(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.machine = Machine.objects.create(code="S3", nom="Machine S3")
        cls.n7 = Produit.objects.create(reference="N7", designation="Produit N7")
        cls.n9 = Produit.objects.create(reference="N9", designation="Produit N9")
        ProduitMachine.objects.create(produit=cls.n7, machine=cls.machine, cadence=200)
        ProduitMachine.objects.create(produit=cls.n9, machine=cls.machine, cadence=150)
        cls.matin = creneau("Matin")
        cls.e1 = Equipe.objects.get(code="E1")
        cls.matricule = Matricule.objects.create(numero="0534")

    def creer_shift(self, debut=None, duree=timedelta(hours=8)):
        debut = debut or a(LUNDI, 5)
        return Shift.objects.create(machine=self.machine, jour=timezone.localtime(debut).date(), creneau=self.matin,
                                    equipe=self.e1, matricule=self.matricule, debut=debut, fin=debut + duree)

    def shift_en_cours(self):
        return self.creer_shift((timezone.now() - timedelta(hours=2)).replace(second=0, microsecond=0))

    def cause(self, libelle):
        return CauseArret.objects.get(libelle=libelle)


class BilanTests(FabriqueShift):
    def test_exemple_de_la_spec(self):
        """480 min, pause 20 et panne 60 : 400 min de marche ; N7 720/800, N9 380/400, shift 1 100/1 200."""
        shift = self.creer_shift()
        n7 = Troncon.objects.create(shift=shift, produit=self.n7, debut=a(LUNDI, 5), compteur_fin=720, cadence=200)
        n9 = Troncon.objects.create(shift=shift, produit=self.n9, debut=a(LUNDI, 9, 20), compteur_debut=100,
                                    compteur_fin=480, cadence=150)
        Arret.objects.create(shift=shift, troncon=n7, cause=self.cause("Pause"), duree=20)
        Arret.objects.create(shift=shift, troncon=n9, cause=self.cause("Panne machine"), duree=60)
        bilan = calculs.bilan_du_shift(shift)
        self.assertEqual((bilan.theorique, bilan.arrets, bilan.marche, bilan.quantite), (480, 80, 400, 1100))
        premier, second = bilan.troncons
        self.assertEqual((premier.duree, premier.marche, premier.attendue, premier.rendement),
                         (260, 240, Decimal(800), Decimal("0.9")))
        self.assertEqual((second.duree, second.marche, second.quantite, second.rendement),
                         (220, 160, 380, Decimal("0.95")))
        self.assertAlmostEqual(float(bilan.rendement), 1100 / 1200)

    def test_quantite_manquante_ou_sans_cadence(self):
        shift = self.creer_shift()
        Troncon.objects.create(shift=shift, produit=self.n7, debut=a(LUNDI, 5), cadence=None, compteur_fin=500)
        Troncon.objects.create(shift=shift, produit=self.n9, debut=a(LUNDI, 10), cadence=150)
        bilan = calculs.bilan_du_shift(shift)
        self.assertEqual(bilan.quantite, 500)
        self.assertIsNone(bilan.rendement)


class SaisieTests(FabriqueShift):
    def setUp(self):
        self.operateur = creer_utilisateur("operateur", roles.OPERATEUR)
        self.client.force_login(self.operateur)

    def test_ouvrir_un_shift_demande_l_equipe(self):
        donnees = {"machine": self.machine.pk, "jour": "2026-10-19", "creneau": self.matin.pk, "matricule": "534"}
        reponse = self.client.post("/production/?machine=S3", donnees)
        self.assertContains(reponse, "Ce champ est obligatoire")
        reponse = self.client.post("/production/?machine=S3", {**donnees, "equipe": self.e1.pk})
        shift = Shift.objects.get()
        self.assertRedirects(reponse, f"/production/shift/{shift.pk}/")
        self.assertEqual((shift.debut, shift.fin, shift.matricule.numero), (a(LUNDI, 5), a(LUNDI, 13), "0534"))
        reponse = self.client.post("/production/?machine=S3", {**donnees, "equipe": self.e1.pk})
        self.assertContains(reponse, "est déjà ouvert")

    def test_l_equipe_n_est_jamais_choisie_d_office(self):
        reponse = self.client.get("/production/?machine=S3")
        self.assertContains(reponse, '<option value="" selected>— Choisir —</option>', html=True)

    def test_produits_changement_et_quantites(self):
        shift = self.shift_en_cours()
        debut = timezone.localtime(shift.debut)
        adresse = f"/production/shift/{shift.pk}/"
        self.client.post(adresse, {"action": "produit", "produit": self.n7.pk, "heure": f"{debut:%H:%M}"})
        n7 = shift.troncons.get()
        self.assertEqual((n7.cadence, n7.compteur_debut), (Decimal("200.0"), 0))
        # Le changement de produit demande le compteur de fin du produit précédent.
        heure_changement = f"{debut + timedelta(hours=1):%H:%M}"
        reponse = self.client.post(adresse, {"action": "produit", "produit": self.n9.pk, "heure": heure_changement})
        self.assertContains(reponse, "Ce champ est obligatoire")
        self.client.post(adresse, {"action": "produit", "produit": self.n9.pk, "heure": heure_changement,
                                   "compteur_fin_precedent": 180, "compteur_debut": 50})
        n7.refresh_from_db()
        self.assertEqual(n7.quantite, 180)
        n9 = shift.troncons.get(produit=self.n9)
        self.client.post(adresse, {"action": "quantite", "troncon": n9.pk, "compteur_debut": 50, "compteur_fin": 30})
        n9.refresh_from_db()
        self.assertIsNone(n9.compteur_fin)  # fin plus petite que le début : refusé
        self.client.post(adresse, {"action": "quantite", "troncon": n9.pk, "compteur_debut": 50, "compteur_fin": 200})
        n9.refresh_from_db()
        self.assertEqual(n9.quantite, 150)
        self.assertFalse(shift.corrections.exists())

    def test_heure_hors_du_shift_refusee(self):
        shift = self.shift_en_cours()
        heure = f"{timezone.localtime(shift.debut) - timedelta(hours=1):%H:%M}"
        reponse = self.client.post(f"/production/shift/{shift.pk}/",
                                   {"action": "produit", "produit": self.n7.pk, "heure": heure})
        self.assertContains(reponse, "L&#x27;heure doit être dans le shift")
        self.assertFalse(shift.troncons.exists())

    def test_arrets(self):
        shift = self.shift_en_cours()
        adresse = f"/production/shift/{shift.pk}/"
        reponse = self.client.post(adresse, {"action": "arret", "cause": self.cause("Autre").pk, "duree": 10})
        self.assertContains(reponse, "Précisez la cause")
        self.client.post(adresse, {"action": "arret", "cause": self.cause("Pause").pk, "duree": 20})
        reponse = self.client.post(adresse, {"action": "arret", "cause": self.cause("Panne machine").pk, "duree": 470})
        self.assertContains(reponse, "490 min d&#x27;arrêt pour 480 min")
        arret = shift.arrets.get()
        self.client.post(adresse, {"action": "retirer_arret", "arret": arret.pk})
        self.assertFalse(shift.arrets.exists())

    def test_shift_termine_corrige_par_methode_avec_historique(self):
        shift = self.creer_shift()
        adresse = f"/production/shift/{shift.pk}/"
        donnees = {"action": "arret", "cause": self.cause("Pause").pk, "duree": 20}
        self.assertEqual(self.client.post(adresse, donnees).status_code, 403)
        self.assertNotContains(self.client.get(adresse), "Ajouter l&#x27;arrêt")
        self.client.force_login(creer_utilisateur("regleur", roles.REGLEUR))
        self.assertEqual(self.client.post(adresse, donnees).status_code, 403)
        methodiste = creer_utilisateur("methode", roles.METHODE)
        self.client.force_login(methodiste)
        self.client.post(adresse, donnees)
        correction = Correction.objects.get()
        self.assertEqual((correction.par, correction.description), (methodiste, "Arrêt ajouté : Pause 20 min."))
        self.assertContains(self.client.get(adresse), "Historique des corrections")

    def test_raccourcis_vers_le_stock_et_les_remontees(self):
        reponse = self.client.get(f"/production/shift/{self.shift_en_cours().pk}/")
        self.assertContains(reponse, 'href="/articles/"')
        self.assertContains(reponse, 'href="/remontees/?machine=S3"')


class RapportTests(FabriqueShift):
    def setUp(self):
        self.client.force_login(creer_utilisateur("methode", roles.METHODE))
        hier = timezone.localdate() - timedelta(days=1)
        self.shift_hier = self.creer_shift(a(hier, 5))
        troncon = Troncon.objects.create(shift=self.shift_hier, produit=self.n7, debut=self.shift_hier.debut,
                                         compteur_fin=1000, cadence=200)
        Arret.objects.create(shift=self.shift_hier, troncon=troncon, cause=self.cause("Pause"), duree=20)
        Arret.objects.create(shift=self.shift_hier, troncon=troncon, cause=self.cause("Panne machine"), duree=60)
        Arret.objects.create(shift=self.shift_hier, troncon=troncon, cause=self.cause("Panne machine"), duree=10,
                             commentaire="Capteur de porte")
        self.shift_ancien = self.creer_shift(a(hier - timedelta(days=10), 5))
        Troncon.objects.create(shift=self.shift_ancien, produit=self.n9, debut=self.shift_ancien.debut,
                               compteur_fin=900, cadence=150)

    def test_s_ouvre_sur_la_veille(self):
        reponse = self.client.get("/production/rapport/")
        self.assertEqual(reponse.context["shifts"], [self.shift_hier])
        self.assertEqual([ligne.cause.libelle for ligne in reponse.context["pareto"]], ["Panne machine", "Pause"])
        self.assertEqual(reponse.context["pareto"][0].minutes, 70)
        self.assertEqual(reponse.context["total"].quantite, 1000)
        self.assertContains(reponse, "Capteur de porte")

    def test_periode_et_filtres(self):
        hier = timezone.localdate() - timedelta(days=1)
        periode = {"du": (hier - timedelta(days=29)).isoformat(), "au": hier.isoformat()}
        reponse = self.client.get("/production/rapport/", periode)
        self.assertEqual(len(reponse.context["shifts"]), 2)
        reponse = self.client.get("/production/rapport/", {**periode, "produit": self.n9.pk})
        self.assertEqual(reponse.context["shifts"], [self.shift_ancien])
        reponse = self.client.get("/production/rapport/", {**periode, "cause": self.cause("Panne machine").pk})
        self.assertEqual(reponse.context["shifts"], [self.shift_hier])

    def test_operateur_voit_seulement_la_veille(self):
        self.client.force_login(creer_utilisateur("operateur", roles.OPERATEUR))
        hier = timezone.localdate() - timedelta(days=1)
        reponse = self.client.get("/production/rapport/", {"du": (hier - timedelta(days=29)).isoformat()})
        self.assertEqual(reponse.context["shifts"], [self.shift_hier])
        self.assertContains(reponse, "Votre profil voit le rapport de la veille")

    def test_export_excel(self):
        reponse = self.client.get("/production/rapport/export.xlsx")
        classeur = load_workbook(BytesIO(reponse.content))
        self.assertEqual(classeur.sheetnames, ["Shifts", "Arrêts"])
        shifts = list(classeur["Shifts"].iter_rows(values_only=True))
        self.assertEqual(shifts[1][3:9], ("S3", "N7", 480, 90, 390, 1000))
        self.assertEqual(len(list(classeur["Arrêts"].iter_rows())), 4)

    def test_sans_droit_pas_de_rapport(self):
        intendante = creer_utilisateur("intendante", roles.INTENDANT)
        groupe = intendante.groups.get()
        groupe.droits.filter(code__startswith="production_").delete()
        self.client.force_login(intendante)
        self.assertEqual(self.client.get("/production/rapport/").status_code, 403)
        self.assertEqual(self.client.get(f"/production/shift/{self.shift_hier.pk}/").status_code, 403)
        self.assertEqual(self.client.get("/production/").status_code, 403)


class DroitsTests(TestCase):
    def test_droits_de_depart_choisis_par_noel(self):
        attendus = {
            roles.METHODE: {"production_saisir", "production_consulter", "production_corriger", "production_parametrer"},
            roles.INTENDANT: {"production_saisir", "production_consulter", "production_parametrer"},
            roles.REGLEUR: {"production_saisir", "production_consulter", "production_parametrer"},
            roles.OPERATEUR: {"production_saisir"},
        }
        for profil, codes in attendus.items():
            utilisateur = creer_utilisateur(f"compte-{len(profil)}-{profil[:3]}", profil)
            self.assertEqual({code for code in roles.droits(utilisateur) if code.startswith("production_")}, codes,
                             profil)


def classeur_historique(lignes):
    """Comme l'ancien classeur : un titre, les filtres, puis l'en-tête (titre « Durée » coupé) et les lignes."""
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Datas"
    feuille.append(["RAPPORT DE PRODUCTION"])
    feuille.append([])
    feuille.append(["Machine", "Jour", "Equipe", "Type arrêt / perte", "Durée (r", "Commentaire"])
    for ligne in lignes:
        feuille.append(ligne)
    contenu = BytesIO()
    classeur.save(contenu)
    contenu.seek(0)
    return contenu


class ImportHistoriqueTests(TestCase):
    def setUp(self):
        for code in ("S1", "S7"):
            Machine.objects.create(code=code, nom=f"Machine {code}")
        self.lignes = [
            ["S7", datetime(2026, 10, 19, 12, 19), "E1", "Organisation", 20, "demarrage lundi"],
            ["S7", "19/10/26 12:20", "E1", "Outillages", 20, "vgl"],
            ["S1", datetime(2026, 10, 19, 19, 45), "E2", "Mensuelle AIC", 90, "Mensuelle + AIC 1"],
            ["S1", datetime(2026, 10, 20, 2, 10), "E3", "Pause", 20, None],
            ["Bande", datetime(2026, 10, 19, 12, 30), "E1", "Pause", 20, None],
        ]

    def test_import_puis_reimport_sans_doublon(self):
        bilan = importer_historique(classeur_historique(self.lignes))
        self.assertEqual((bilan.arrets_crees, bilan.shifts_crees, bilan.doublons), (4, 3, 0))
        self.assertEqual(bilan.causes_creees, ["Mensuelle AIC"])
        self.assertEqual(bilan.ignorees, [(8, "machine inconnue « Bande »")])
        s7 = Shift.objects.get(machine__code="S7")
        self.assertEqual((s7.creneau.nom, s7.equipe.code, s7.arrets.count()), ("Matin", "E1", 2))
        nuit = Shift.objects.get(creneau__nom="Nuit")
        self.assertEqual((nuit.jour, nuit.equipe.code), (date(2026, 10, 19), "E3"))
        self.assertIsNone(nuit.bilan().rendement)
        bilan = importer_historique(classeur_historique(self.lignes))
        self.assertEqual((bilan.arrets_crees, bilan.doublons), (0, 4))
        self.assertEqual(Arret.objects.count(), 4)
